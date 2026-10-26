import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_docx():
    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Palette: Deep Blue (#1B365D), Accent Teal (#008080), Charcoal (#2C3E50), Soft Gray (#F4F6F9)
    COLOR_PRIMARY = RGBColor(27, 54, 93)     # #1B365D
    COLOR_SECONDARY = RGBColor(0, 128, 128)  # #008080
    COLOR_BODY = RGBColor(44, 62, 80)        # #2C3E50
    COLOR_MUTED = RGBColor(100, 110, 120)

    # Document Header / Badge
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_badge = p_badge.add_run("INAMIGOS FOUNDATION INTERNSHIP PROGRAM • TASK 1")
    run_badge.font.name = "Arial"
    run_badge.font.size = Pt(9.5)
    run_badge.font.bold = True
    run_badge.font.color.rgb = COLOR_SECONDARY

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Understanding Artificial Intelligence & Machine Learning")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("A Comprehensive Beginner's Learning Guide with Real-World Examples")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_MUTED

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    run_div = p_div.add_run("―" * 58)
    run_div.font.color.rgb = RGBColor(200, 210, 220)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Executive Overview Box
    table_box = doc.add_table(rows=1, cols=1)
    table_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_box.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)

    p_box = cell.paragraphs[0]
    r_box_title = p_box.add_run("🎯 Learning Objective\n")
    r_box_title.bold = True
    r_box_title.font.name = "Arial"
    r_box_title.font.size = Pt(11)
    r_box_title.font.color.rgb = COLOR_PRIMARY

    r_box_text = p_box.add_run(
        "This learning guide breaks down the core concepts that power modern intelligent technology—Artificial Intelligence, "
        "Machine Learning, Deep Learning, and the three primary learning paradigms (Supervised, Unsupervised, and Reinforcement Learning). "
        "Each concept is explained in clear, intuitive language paired with an accessible, real-world example."
    )
    r_box_text.font.name = "Arial"
    r_box_text.font.size = Pt(10)
    r_box_text.font.color.rgb = COLOR_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Core Sections Data
    concepts = [
        {
            "num": "01",
            "title": "Artificial Intelligence (AI)",
            "subtitle": "The Broad Vision of Smart Machines",
            "concept_def": (
                "Artificial Intelligence (AI) is the overarching branch of computer science dedicated to building machines "
                "and systems capable of performing tasks that typically require human intelligence. These tasks include reasoning, "
                "problem-solving, visual perception, decision-making, understanding language, and adapting to new situations."
            ),
            "how_it_works": (
                "Instead of relying on fixed manual labor, AI systems process data, apply logic or learned patterns, and select "
                "the best action to accomplish a specified goal."
            ),
            "example_title": "Real-Life Example: Smart Virtual Assistants (Apple Siri / Google Assistant / Amazon Alexa)",
            "example_desc": (
                "When you say 'Hey Google, what's the weather today in Mumbai?', the AI system recognizes your spoken words (speech recognition), "
                "understands your intent (natural language understanding), retrieves weather forecast data, and speaks back a natural reply."
            )
        },
        {
            "num": "02",
            "title": "Machine Learning (ML)",
            "subtitle": "Learning from Experience and Data",
            "concept_def": (
                "Machine Learning (ML) is a vital subset of AI where computers are not explicitly programmed with rigid rules. "
                "Instead, algorithms are fed historical data and learn to detect patterns, improve their own accuracy over time, "
                "and make reliable predictions on new, unseen data."
            ),
            "how_it_works": (
                "Traditional programming requires human developers to write every rule ('If X, do Y'). Machine Learning flips this model: "
                "we feed input data and output answers into the algorithm, and the computer formulates the rules itself."
            ),
            "example_title": "Real-Life Example: Email Spam Detection & Music Recommendations (Spotify / Netflix)",
            "example_desc": (
                "Gmail inspects millions of incoming emails. By analyzing words, sender reputation, and user clicks ('Report Spam'), "
                "the ML model automatically classifies future emails as spam or inbox without anyone writing rules for every spam message."
            )
        },
        {
            "num": "03",
            "title": "Deep Learning (DL)",
            "subtitle": "Multi-Layered Brain-Inspired Neural Networks",
            "concept_def": (
                "Deep Learning (DL) is an advanced subset of Machine Learning powered by Artificial Neural Networks with multiple hidden layers "
                "(hence 'deep'). It mimics the structural information processing of neurons in the human brain, allowing systems to process massive, "
                "unstructured datasets like images, audio, video, and raw text with high precision."
            ),
            "how_it_works": (
                "Unlike standard ML where humans often have to manually highlight important features, Deep Learning automatically discovers intricate "
                "features across hierarchical layers (e.g., edges → shapes → eyes/noses → full faces)."
            ),
            "example_title": "Real-Life Example: Smartphone Face Recognition (Face ID) & Autonomous Self-Driving Cars",
            "example_desc": (
                "When unlocking your phone with Face ID, a deep neural network maps thousands of 3D infrared facial data points. "
                "It reliably identifies your face even if you wear glasses, change your hairstyle, or stand in low lighting."
            )
        },
        {
            "num": "04",
            "title": "Supervised Learning",
            "subtitle": "Learning with a Teacher (Labeled Data)",
            "concept_def": (
                "Supervised Learning is the most common machine learning approach where the algorithm is trained on a labeled dataset. "
                "Every training example contains both the input features (questions) and the correct target label (answers). The model learns "
                "the relationship between inputs and outputs so it can predict labels for new, unlabeled data."
            ),
            "how_it_works": (
                "Think of a student studying with flashcards containing questions on the front and correct answers on the back. "
                "The teacher corrects errors until the student consistently produces the right answer."
            ),
            "example_title": "Real-Life Example: House Price Prediction & Medical Image Diagnostics",
            "example_desc": (
                "A real estate model is trained on thousands of home sales with labeled data: square footage, number of bedrooms, location, "
                "and final sale price. Once trained, you can input a newly listed house's specs, and the model predicts its market value accurately."
            )
        },
        {
            "num": "05",
            "title": "Unsupervised Learning",
            "subtitle": "Discovering Hidden Patterns (Unlabeled Data)",
            "concept_def": (
                "Unsupervised Learning occurs when algorithms are provided with raw, unlabeled data without any predefined answers or target outcomes. "
                "The model's objective is to explore the data independently, uncovering hidden structures, clusters, similarities, and anomalies."
            ),
            "how_it_works": (
                "Imagine giving a child a mixed basket of fruits they have never seen before without names. The child naturally groups "
                "the round red ones together (apples), the long yellow ones together (bananas), and the small purple ones together (grapes)."
            ),
            "example_title": "Real-Life Example: Customer Segmentation in E-Commerce & Bank Fraud Detection",
            "example_desc": (
                "Amazon or Flipkart analyzes millions of shoppers with unsupervised clustering. The algorithm automatically groups shoppers into segments "
                "(e.g., 'bargain hunters', 'gadget enthusiasts', 'weekend impulse buyers'), allowing marketing teams to target each group uniquely."
            )
        },
        {
            "num": "06",
            "title": "Reinforcement Learning (RL)",
            "subtitle": "Learning by Trial, Error, and Reward",
            "concept_def": (
                "Reinforcement Learning (RL) is a goal-oriented learning framework where an autonomous agent learns to make optimal decisions "
                "by interacting directly with a dynamic environment. The agent receives feedback in the form of numerical rewards for desirable actions "
                "and penalties for mistakes, continuously updating its strategy to maximize long-term cumulative reward."
            ),
            "how_it_works": (
                "Similar to training a pet dog: when the dog obeys a command ('sit'), it gets a treat (reward); when it misbehaves, it receives no reward (penalty). "
                "Over repeated trials, the dog learns which behaviors bring the maximum treats."
            ),
            "example_title": "Real-Life Example: Self-Driving Vehicles & Game AI (AlphaGo / Video Game Bots)",
            "example_desc": (
                "Google DeepMind's AlphaGo learned to master the game of Go by playing millions of games against itself. "
                "By receiving positive rewards for winning positions and penalties for moves leading to losses, it developed strategies superior to world human champions."
            )
        }
    ]

    for item in concepts:
        p_h = doc.add_paragraph()
        p_h.paragraph_format.space_before = Pt(14)
        p_h.paragraph_format.space_after = Pt(2)
        
        r_num = p_h.add_run(f"[{item['num']}]  ")
        r_num.bold = True
        r_num.font.name = "Arial"
        r_num.font.size = Pt(13)
        r_num.font.color.rgb = COLOR_SECONDARY

        r_htitle = p_h.add_run(item['title'])
        r_htitle.bold = True
        r_htitle.font.name = "Arial"
        r_htitle.font.size = Pt(14)
        r_htitle.font.color.rgb = COLOR_PRIMARY

        p_subt = doc.add_paragraph()
        p_subt.paragraph_format.space_after = Pt(6)
        r_subt = p_subt.add_run(item['subtitle'])
        r_subt.italic = True
        r_subt.font.name = "Arial"
        r_subt.font.size = Pt(10)
        r_subt.font.color.rgb = COLOR_MUTED

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_after = Pt(4)
        r_desc_b = p_desc.add_run("• What it means: ")
        r_desc_b.bold = True
        r_desc_b.font.name = "Arial"
        r_desc_b.font.size = Pt(10.5)
        r_desc_b.font.color.rgb = COLOR_BODY
        r_desc = p_desc.add_run(item['concept_def'])
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(10.5)
        r_desc.font.color.rgb = COLOR_BODY

        p_how = doc.add_paragraph()
        p_how.paragraph_format.space_after = Pt(6)
        r_how_b = p_how.add_run("• Intuition / How it works: ")
        r_how_b.bold = True
        r_how_b.font.name = "Arial"
        r_how_b.font.size = Pt(10.5)
        r_how_b.font.color.rgb = COLOR_BODY
        r_how = p_how.add_run(item['how_it_works'])
        r_how.font.name = "Arial"
        r_how.font.size = Pt(10.5)
        r_how.font.color.rgb = COLOR_BODY

        # Callout box for Real-Life Example
        tbl_ex = doc.add_table(rows=1, cols=1)
        tbl_ex.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_ex = tbl_ex.cell(0, 0)
        set_cell_background(cell_ex, "EBF5FB") # Soft blue
        set_cell_margins(cell_ex, top=100, bottom=100, left=150, right=150)
        
        p_ex = cell_ex.paragraphs[0]
        r_ex_title = p_ex.add_run(f"🌟 {item['example_title']}\n")
        r_ex_title.bold = True
        r_ex_title.font.name = "Arial"
        r_ex_title.font.size = Pt(10)
        r_ex_title.font.color.rgb = COLOR_PRIMARY

        r_ex_text = p_ex.add_run(item['example_desc'])
        r_ex_text.font.name = "Arial"
        r_ex_text.font.size = Pt(9.5)
        r_ex_text.font.color.rgb = COLOR_BODY

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Page Break for Summary & Comparison
    doc.add_page_break()

    # Section 7: Quick Reference Summary Table
    p_tbl_header = doc.add_paragraph()
    p_tbl_header.paragraph_format.space_before = Pt(10)
    p_tbl_header.paragraph_format.space_after = Pt(8)
    r_th = p_tbl_header.add_run("Quick Reference: Core AI & ML Comparison Table")
    r_th.bold = True
    r_th.font.name = "Arial"
    r_th.font.size = Pt(14)
    r_th.font.color.rgb = COLOR_PRIMARY

    # Table creation: 4 columns
    comp_data = [
        ["Concept", "Core Question / Focus", "Key Characteristic", "Real-Life Example"],
        ["Artificial Intelligence", "Can machines act smart?", "Broad umbrella; mimics human thinking", "Siri, Alexa, Smart GPS"],
        ["Machine Learning", "How do machines learn?", "Learns patterns from data without explicit code", "Spam filter, Netflix picks"],
        ["Deep Learning", "Can machines learn complex tasks?", "Multi-layered neural networks (like brain)", "Face ID, Autonomous cars"],
        ["Supervised Learning", "How to predict known targets?", "Trained on labeled inputs and outputs", "House price prediction"],
        ["Unsupervised Learning", "How to discover hidden patterns?", "Trained on raw, unlabeled data", "Customer clustering"],
        ["Reinforcement Learning", "How to take optimal actions?", "Agent learns via rewards and penalties", "AlphaGo, Game bots, Robotics"]
    ]

    tbl_comp = doc.add_table(rows=len(comp_data), cols=4)
    tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_idx, row in enumerate(tbl_comp.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.text = comp_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.runs[0]
            r.font.name = "Arial"
            if row_idx == 0:
                set_cell_background(cell, "1B365D")
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_BODY
                if col_idx == 0:
                    r.font.bold = True
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)

    # Conceptual Relationship Hierarchy
    p_rel = doc.add_paragraph()
    p_rel.paragraph_format.space_before = Pt(18)
    p_rel.paragraph_format.space_after = Pt(6)
    r_rel_h = p_rel.add_run("Conceptual Hierarchy: How Everything Fits Together")
    r_rel_h.bold = True
    r_rel_h.font.name = "Arial"
    r_rel_h.font.size = Pt(13)
    r_rel_h.font.color.rgb = COLOR_PRIMARY

    p_rel_body = doc.add_paragraph()
    p_rel_body.paragraph_format.space_after = Pt(10)
    p_rel_body.add_run(
        "• Artificial Intelligence is the broad universe (any intelligent machine).\n"
        "  └── Machine Learning is the driving engine inside AI (learning from data).\n"
        "        └── Deep Learning is the specialized core inside ML (multi-layered neural nets).\n"
        "• The 3 Learning Paradigms describe HOW machines learn:\n"
        "  1. Supervised Learning: Learning with guided supervision (labeled data).\n"
        "  2. Unsupervised Learning: Self-directed discovery (unlabeled data).\n"
        "  3. Reinforcement Learning: Learning through environmental feedback (rewards/penalties)."
    )
    for r in p_rel_body.runs:
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_BODY

    # Footer note
    p_ft = doc.add_paragraph()
    p_ft.paragraph_format.space_before = Pt(20)
    p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ft = p_ft.add_run("Prepared for InAmigos Foundation • Author: Yamuna • Task 1 Submission")
    r_ft.font.name = "Arial"
    r_ft.font.size = Pt(9)
    r_ft.font.italic = True
    r_ft.font.color.rgb = COLOR_MUTED

    output_path = r"d:\Yamuna akka files\Understanding_AI_and_Machine_Learning_Learning_Guide.docx"
    doc.save(output_path)
    print(f"Successfully generated Word document: {output_path}")

if __name__ == "__main__":
    create_docx()
