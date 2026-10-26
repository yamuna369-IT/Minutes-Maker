from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_bullets(doc: Document, values: list[str], empty_text: str) -> None:
    if not values:
        doc.add_paragraph(empty_text)
        return

    for value in values:
        doc.add_paragraph(str(value), style="List Bullet")


def export_docx(data: dict, out_path: str) -> str:
    doc = Document()

    title = doc.add_heading(data["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(f"Generated: {data['date']} | Language: {data['language'].upper()}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(data["summary"])

    doc.add_heading("Key Decisions", level=1)
    _add_bullets(doc, data["decisions"], "No explicit decisions were detected.")

    doc.add_heading("Action Items", level=1)
    _add_bullets(doc, data["action_items"], "No action items were detected.")

    doc.add_heading("Keywords", level=1)
    doc.add_paragraph(", ".join(data["keywords"]) if data["keywords"] else "No keywords detected.")

    doc.add_heading("Diarized Transcript", level=1)
    for segment in data["diarized_transcript"]:
        text = segment.get("text", "").strip()
        if not text:
            continue
        doc.add_paragraph(
            f"{segment['speaker']} [{segment['start']:.1f}-{segment['end']:.1f}]",
            style="Intense Quote",
        )
        doc.add_paragraph(text)

    doc.save(out_path)
    return out_path
